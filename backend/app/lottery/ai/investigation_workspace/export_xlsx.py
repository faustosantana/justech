"""Export InvestigationAsset to XLSX (Resumen + Resultados)."""

from __future__ import annotations

import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.config import settings
from app.lottery.ai.investigation_workspace.schemas import InvestigationAsset
from app.lottery.ai.official_lottery_scope import OFFICIAL_LOTTERY_SCOPE, scope_label_es


def _exports_dir() -> Path:
    root = Path(settings.lottery_exports_path).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def _safe_cell(value: Any) -> str | int | float | None:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    s = str(value)
    if s and s[0] in "=+@-\t\r":
        return "'" + s
    return s


def export_asset_xlsx(asset: InvestigationAsset) -> tuple[Path, str, bytes]:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Resumen"
    bold = Font(bold=True)
    meta = [
        ("Título", asset.title),
        ("Fecha de exportación", datetime.now(timezone.utc).isoformat()),
        ("Asset ID", asset.asset_id),
        ("Investigation ID", asset.investigation_id or ""),
        ("Sujetos", ", ".join(asset.subjects)),
        ("Relación", asset.relation or ""),
        ("Alcance", scope_label_es()),
        ("Loterías oficiales", ", ".join(OFFICIAL_LOTTERY_SCOPE)),
        ("Filtros", str(asset.filters.model_dump(exclude_none=True))),
        ("Orden", str(asset.sort.model_dump())),
        ("Filas exportadas", asset.row_count),
        ("Filas fuente", len(asset.source_rows)),
    ]
    ws_sum["A1"] = "Resumen de exportación"
    ws_sum["A1"].font = bold
    for i, (k, v) in enumerate(meta, start=3):
        ws_sum.cell(i, 1, k).font = bold
        ws_sum.cell(i, 2, _safe_cell(v))

    ws = wb.create_sheet("Resultados")
    cols = list(asset.columns)
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(1, c, name)
        cell.font = bold
    for r_i, row in enumerate(asset.rows, start=2):
        for c_i, name in enumerate(cols, start=1):
            ws.cell(r_i, c_i, _safe_cell(row.get(name)))

    bio_path = _exports_dir()
    storage = f"ws_{asset.asset_id}_{uuid.uuid4().hex[:10]}.xlsx"
    path = bio_path / storage
    wb.save(path)
    content = path.read_bytes()
    slug_bits = "-".join(asset.subjects) or "asset"
    slug_bits = re.sub(r"[^a-zA-Z0-9._-]+", "-", slug_bits)
    filename = f"workspace-{slug_bits}-{asset.asset_id[:8]}.xlsx"
    return path, filename, content
