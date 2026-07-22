"""Vista previa PDF — LibreOffice (principal) o ReportLab (fallback)."""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.document_autofill.libreoffice_converter import (
    docx_bytes_to_pdf_libreoffice,
    is_libreoffice_available,
)
from app.services.document_autofill.pdf_field_writer import fill_pdf_acroform

logger = logging.getLogger(__name__)


def docx_bytes_to_pdf_reportlab(docx_bytes: bytes, title: str = "Documento") -> bytes:
    """Fallback — render aproximado cuando LibreOffice no está disponible."""
    doc = Document(BytesIO(docx_bytes))
    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title=title,
    )
    styles = getSampleStyleSheet()
    story: list = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            story.append(Spacer(1, 6))
            continue
        style = styles["Heading2"] if para.runs and para.runs[0].bold else styles["Normal"]
        safe = text.replace("&", "&amp;").replace("<", "&lt;")
        if "[[PENDIENTE]]" in safe:
            safe = f'<font color="red">{safe}</font>'
        story.append(Paragraph(safe, style))
        story.append(Spacer(1, 4))

    for table in doc.tables:
        data = [[cell.text for cell in row.cells] for row in table.rows]
        if not data:
            continue
        tbl = Table(data, repeatRows=1)
        tbl.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(tbl)
        story.append(Spacer(1, 12))

    pdf.build(story)
    buffer.seek(0)
    return buffer.read()


def docx_bytes_to_pdf(docx_bytes: bytes, title: str = "Documento") -> tuple[bytes, str]:
    """Convierte DOCX a PDF. Retorna (bytes, engine) donde engine es libreoffice|reportlab."""
    if is_libreoffice_available():
        try:
            return docx_bytes_to_pdf_libreoffice(docx_bytes), "libreoffice"
        except Exception as exc:
            logger.warning("LibreOffice conversion failed, using ReportLab fallback: %s", exc)
    return docx_bytes_to_pdf_reportlab(docx_bytes, title=title), "reportlab"


def template_to_preview_pdf(
    *,
    template_bytes: bytes,
    template_format: str,
    filled_docx_bytes: bytes | None,
    values: dict[str, str],
    title: str,
) -> tuple[bytes, str]:
    if template_format == "pdf":
        try:
            reader = PdfReader(BytesIO(template_bytes))
            if reader.get_fields():
                return fill_pdf_acroform(template_bytes, values), "pypdf_acroform"
        except Exception:
            pass
    if filled_docx_bytes:
        return docx_bytes_to_pdf(filled_docx_bytes, title=title)
    return docx_bytes_to_pdf(template_bytes, title=title)


def write_pdf(path: Path, pdf_bytes: bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(pdf_bytes)
    return path
