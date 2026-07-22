"""Genera plantillas DOCX base con placeholders (solo si no existen)."""

from __future__ import annotations

from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _add_field_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (label, key) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = f"{{{{ {key} }}}}"


def _write_carta_presentacion(path: Path) -> None:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("CARTA DE PRESENTACIÓN DE OFERTA")
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph()
    doc.add_paragraph("Fecha: {{ fecha }}")
    doc.add_paragraph()
    doc.add_paragraph("Señores")
    doc.add_paragraph("{{ entidad_contratante }}")
    doc.add_paragraph()
    doc.add_paragraph(
        "Por medio de la presente, {{ razon_social }}, RNC {{ rnc }}, con domicilio en "
        "{{ direccion }}, representada por {{ representante_legal }}, presenta formalmente "
        "su oferta al proceso de contratación {{ proceso_dgcp }}."
    )
    doc.add_paragraph()
    doc.add_paragraph("Objeto: {{ objeto_proceso }}")
    doc.add_paragraph("Monto ofertado: {{ monto }}")
    doc.add_paragraph("Plazo: {{ plazo_entrega }}")
    doc.add_paragraph("Condiciones: {{ condiciones }}")
    doc.add_paragraph()
    doc.add_paragraph("Contacto: {{ telefono }} · {{ correo }}")
    doc.add_paragraph()
    doc.add_paragraph("Atentamente,")
    doc.add_paragraph()
    doc.add_paragraph("_______________________________")
    doc.add_paragraph("{{ representante_legal }}")
    doc.add_paragraph("{{ razon_social }}")
    doc.add_paragraph()
    doc.add_paragraph("[Espacio para firma y sello]")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _write_sncc(path: Path, code: str, title: str, rows: list[tuple[str, str]]) -> None:
    doc = Document()
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(f"{code} — {title}")
    r.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph(f"Proceso DGCP: {{{{ proceso_dgcp }}}}")
    doc.add_paragraph(f"Entidad: {{{{ entidad_contratante }}}}")
    doc.add_paragraph()
    _add_field_table(doc, rows)
    doc.add_paragraph()
    doc.add_paragraph("Declaro bajo fe de juramento que la información es verídica.")
    doc.add_paragraph()
    doc.add_paragraph("Fecha: {{ fecha }}")
    doc.add_paragraph()
    doc.add_paragraph("_______________________________")
    doc.add_paragraph("{{ representante_legal }}")
    doc.add_paragraph("[Firma y sello]")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _write_oferta_economica(path: Path) -> None:
    doc = Document()
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("OFERTA ECONÓMICA")
    r.bold = True
    doc.add_paragraph("Proceso: {{ proceso_dgcp }}")
    doc.add_paragraph("Oferente: {{ razon_social }} · RNC {{ rnc }}")
    doc.add_paragraph()
    _add_field_table(
        doc,
        [
            ("Objeto", "objeto_proceso"),
            ("Monto total", "monto"),
            ("Productos / servicios", "productos"),
            ("Plazo de entrega", "plazo_entrega"),
            ("Garantía", "garantia"),
            ("Condiciones de pago", "condiciones"),
            ("Cuenta bancaria", "cuenta_bancaria"),
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph("Fecha: {{ fecha }}")
    doc.add_paragraph("[Firma y sello]")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def ensure_all_templates(root: Path) -> list[Path]:
    created: list[Path] = []
    specs: list[tuple[str, str, Callable[[Path], None]]] = [
        (
            "dgcp/sncc-f034.docx",
            "sncc-f034",
            lambda p: _write_sncc(
                p,
                "SNCC.F.034",
                "Presentación de Oferta",
                [
                    ("Razón social", "razon_social"),
                    ("RNC", "rnc"),
                    ("Representante legal", "representante_legal"),
                    ("Dirección", "direccion"),
                    ("Teléfono", "telefono"),
                    ("Correo", "correo"),
                ],
            ),
        ),
        (
            "dgcp/sncc-f042.docx",
            "sncc-f042",
            lambda p: _write_sncc(
                p,
                "SNCC.F.042",
                "Información del Oferente",
                [
                    ("Razón social", "razon_social"),
                    ("RNC", "rnc"),
                    ("Dirección", "direccion"),
                    ("Representante legal", "representante_legal"),
                    ("Teléfono", "telefono"),
                    ("Correo", "correo"),
                ],
            ),
        ),
        (
            "dgcp/sncc-f033.docx",
            "sncc-f033",
            lambda p: _write_sncc(
                p,
                "SNCC.F.033",
                "Declaración Jurada",
                [
                    ("Razón social", "razon_social"),
                    ("RNC", "rnc"),
                    ("Representante legal", "representante_legal"),
                ],
            ),
        ),
        (
            "dgcp/sncc-f047.docx",
            "sncc-f047",
            lambda p: _write_sncc(
                p,
                "SNCC.F.047",
                "Información Bancaria",
                [
                    ("Razón social", "razon_social"),
                    ("RNC", "rnc"),
                    ("Dirección", "direccion"),
                    ("Cuenta bancaria", "cuenta_bancaria"),
                ],
            ),
        ),
        ("commercial/carta-presentacion.docx", "carta-presentacion", _write_carta_presentacion),
        ("commercial/oferta-economica.docx", "oferta-economica", _write_oferta_economica),
    ]
    for rel, _slug, writer in specs:
        path = root / rel
        if path.is_file():
            continue
        writer(path)
        created.append(path)
    return created
