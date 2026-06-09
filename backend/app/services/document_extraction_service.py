"""Extracción de texto — PDF, DOCX, XLSX, TXT, CSV, RTF, ODT."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field


@dataclass
class ExtractionResult:
    text: str
    pages: list[str] = field(default_factory=list)
    format: str = "other"


class DocumentExtractionService:
    SUPPORTED = frozenset({"pdf", "docx", "xlsx", "txt", "csv", "rtf", "odt"})

    @classmethod
    def detect_format(cls, filename: str, mime_type: str | None = None) -> str:
        lower = filename.lower()
        for ext in ("pdf", "docx", "xlsx", "txt", "csv", "rtf", "odt"):
            if lower.endswith(f".{ext}"):
                return ext
        if mime_type:
            mapping = {
                "application/pdf": "pdf",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
                "text/plain": "txt",
                "text/csv": "csv",
                "application/rtf": "rtf",
            }
            return mapping.get(mime_type, "other")
        return "other"

    def extract(self, content: bytes, *, filename: str, mime_type: str | None = None) -> ExtractionResult:
        fmt = self.detect_format(filename, mime_type)
        if fmt == "txt" or fmt == "csv":
            return self._extract_text(content, fmt)
        if fmt == "pdf":
            return self._extract_pdf(content)
        if fmt == "docx":
            return self._extract_docx(content)
        if fmt == "xlsx":
            return self._extract_xlsx(content)
        if fmt == "rtf":
            return self._extract_rtf(content)
        if fmt == "odt":
            return self._extract_odt(content)
        try:
            text = content.decode("utf-8", errors="ignore")
            return ExtractionResult(text=text.strip(), pages=[text.strip()] if text.strip() else [], format=fmt)
        except Exception:
            return ExtractionResult(text="", pages=[], format=fmt)

    @staticmethod
    def _extract_text(content: bytes, fmt: str) -> ExtractionResult:
        text = content.decode("utf-8", errors="ignore")
        if fmt == "csv":
            reader = csv.reader(io.StringIO(text))
            rows = [" | ".join(row) for row in reader]
            text = "\n".join(rows)
        return ExtractionResult(text=text.strip(), pages=[text.strip()] if text.strip() else [], format=fmt)

    @staticmethod
    def ocr_available() -> bool:
        try:
            import pytesseract  # noqa: F401
            from pdf2image import convert_from_bytes  # noqa: F401

            return True
        except ImportError:
            return False

    @staticmethod
    def _extract_pdf_ocr(content: bytes) -> ExtractionResult:
        try:
            import pytesseract
            from pdf2image import convert_from_bytes

            images = convert_from_bytes(content, dpi=200)
            pages: list[str] = []
            for image in images:
                pages.append(pytesseract.image_to_string(image, lang="spa+eng").strip())
            full = "\n\n".join(p for p in pages if p)
            return ExtractionResult(text=full, pages=pages, format="pdf")
        except Exception:
            return ExtractionResult(text="", pages=[], format="pdf")

    @staticmethod
    def _extract_pdf(content: bytes) -> ExtractionResult:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            pages: list[str] = []
            for page in reader.pages:
                pages.append((page.extract_text() or "").strip())
            full = "\n\n".join(p for p in pages if p)
            if full.strip():
                return ExtractionResult(text=full, pages=pages, format="pdf")
            if DocumentExtractionService.ocr_available():
                ocr_result = DocumentExtractionService._extract_pdf_ocr(content)
                if ocr_result.text.strip():
                    return ocr_result
            return ExtractionResult(text=full, pages=pages, format="pdf")
        except Exception:
            if DocumentExtractionService.ocr_available():
                return DocumentExtractionService._extract_pdf_ocr(content)
            return ExtractionResult(text="", pages=[], format="pdf")

    @staticmethod
    def _extract_docx(content: bytes) -> ExtractionResult:
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            return ExtractionResult(text=text, pages=[text] if text else [], format="docx")
        except Exception:
            return DocumentExtractionService._extract_zip_xml(content, "word/document.xml", "docx")

    @staticmethod
    def _extract_xlsx(content: bytes) -> ExtractionResult:
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            text = "\n".join(lines)
            return ExtractionResult(text=text, pages=[text] if text else [], format="xlsx")
        except Exception:
            return ExtractionResult(text="", pages=[], format="xlsx")

    @staticmethod
    def _extract_rtf(content: bytes) -> ExtractionResult:
        raw = content.decode("latin-1", errors="ignore")
        text = re.sub(r"\\[a-z]+\d* ?", "", raw)
        text = re.sub(r"[{}]", "", text)
        text = re.sub(r"\s+", " ", text).strip()
        return ExtractionResult(text=text, pages=[text] if text else [], format="rtf")

    @staticmethod
    def _extract_odt(content: bytes) -> ExtractionResult:
        return DocumentExtractionService._extract_zip_xml(content, "content.xml", "odt")

    @staticmethod
    def _extract_zip_xml(content: bytes, inner_path: str, fmt: str) -> ExtractionResult:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                raw = zf.read(inner_path).decode("utf-8", errors="ignore")
            text = re.sub(r"<[^>]+>", " ", raw)
            text = re.sub(r"\s+", " ", text).strip()
            return ExtractionResult(text=text, pages=[text] if text else [], format=fmt)
        except Exception:
            return ExtractionResult(text="", pages=[], format=fmt)

    @staticmethod
    def chunk_pages(pages: list[str], *, max_chunk: int = 2000) -> list[tuple[int | None, int, str]]:
        chunks: list[tuple[int | None, int, str]] = []
        idx = 0
        for page_num, page_text in enumerate(pages, start=1):
            if not page_text:
                continue
            if len(page_text) <= max_chunk:
                chunks.append((page_num, idx, page_text))
                idx += 1
                continue
            for i in range(0, len(page_text), max_chunk):
                chunks.append((page_num, idx, page_text[i : i + max_chunk]))
                idx += 1
        if not chunks and pages:
            joined = "\n".join(pages)
            chunks.append((1, 0, joined[:max_chunk]))
        return chunks
