"""Parser de listas de precios XLSX/XLS/CSV — preserva fila original y precios múltiples."""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from app.services.price_classification import classify_product
from app.services.price_normalizer import (
    PRICE_MIN,
    build_search_blob,
    detect_brand,
    detect_supplier_from_path,
    normalize_header,
    parse_int_quantity,
    parse_price,
    parse_ram_gb,
    parse_storage,
)
from app.services.price_sheet_policy import is_auxiliary_sheet, is_commercial_sheet, sheet_skip_reason

PRICE_COLUMN_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("price_final", ("precio final", "your price", "dealer price", "sale price", "net price")),
    ("price_discount", ("precio con descuento", "discount price", "sale price", "discounted")),
    ("price_regular", ("precio regular", "regular price", "p. unitario", "unit price", "list price", "precio")),
    ("price_rebate", ("rebate", "rebaja", "descuento")),
    ("price_cost", ("costo", "cost", "dealer cost")),
]

PREFERRED_PRICE_ORDER = ("price_final", "price_discount", "price_regular", "price_cost")

STOCK_COLUMN_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("stock", ("stock", "oh", "stk", "stcok", "available", "disponible", "existencia", "on hand", "qty", "quantity")),
    ("transit", ("transito", "tránsito", "opo", "in transit", "transito")),
]

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "sku": ("sku", "product id", "codigo", "código", "codigo local", "cod. fabricante", "vpn"),
    "mpn": ("mpn", "part number", "vendor part number", "part number (mpn)"),
    "description": (
        "descripcion", "descripción", "description", "product description",
        "product description local", "descripcion local", "short description",
    ),
    "category": ("categoria", "categoría", "category", "familia", "line", "lob", "line of business"),
    "processor": ("procesador", "processor"),
    "ram": ("ram",),
    "storage": ("storage hard drive", "storage", "almacenamiento", "disco"),
    "display": ("display", "pantalla", "resolucion"),
    "os": ("so", "sistema operativo", "operating system"),
    "warranty": ("warranty", "garantia", "garantía", "base warranty"),
    "model": ("modelo", "model", "familia"),
}


@dataclass
class SheetAudit:
    sheet_name: str
    status: str  # indexed | ignored | empty
    reason: str | None = None
    header_row: int = 0
    columns_mapped: dict[str, str] = field(default_factory=dict)
    rows_read: int = 0
    rows_valid: int = 0
    rows_discarded: int = 0
    discard_reasons: dict[str, int] = field(default_factory=dict)


@dataclass
class ParsedPriceRow:
    sku: str | None = None
    mpn: str | None = None
    model: str | None = None
    description: str | None = None
    category: str | None = None
    product_type: str = "general"
    brand: str | None = None
    processor: str | None = None
    ram_gb: int | None = None
    storage_gb: int | None = None
    storage_type: str | None = None
    display: str | None = None
    operating_system: str | None = None
    price: Decimal | None = None
    preferred_price: Decimal | None = None
    preferred_price_field: str | None = None
    price_regular: Decimal | None = None
    price_rebate: Decimal | None = None
    price_discount: Decimal | None = None
    prices_original: dict[str, Any] = field(default_factory=dict)
    currency: str = "USD"
    stock: int | None = None
    in_transit: int | None = None
    stock_text_original: str | None = None
    warranty: str | None = None
    sheet_name: str | None = None
    row_number: int | None = None
    raw_row_json: dict[str, Any] = field(default_factory=dict)
    raw_columns_json: list[dict[str, Any]] = field(default_factory=list)
    search_blob: str = ""
    is_commercial: bool = True
    excluded_from_laptop: bool = True
    is_cotizable: bool = False
    price_review_status: str = "requires_review"
    classification_label: str = "producto general"
    stock_source_column: str | None = None


@dataclass
class ParseResult:
    rows: list[ParsedPriceRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    sheet_audits: list[SheetAudit] = field(default_factory=list)
    supplier: str | None = None
    manufacturer: str | None = None

    @property
    def laptops_count(self) -> int:
        return sum(1 for r in self.rows if r.product_type == "laptop" and not r.excluded_from_laptop)

    @property
    def with_price_count(self) -> int:
        return sum(1 for r in self.rows if r.preferred_price is not None)

    @property
    def with_stock_count(self) -> int:
        return sum(1 for r in self.rows if r.stock is not None)


class PriceListParser:
    SUPPORTED = frozenset({".xlsx", ".xls", ".csv"})
    SECTION_HEADERS = frozenset({
        "desktops", "monitores", "notebooks", "laptops", "hospitality", "desktop", "monitor",
        "14 pulgadas", "15 pulgadas", "16 pulgadas", "p. comerciales",
    })

    def parse_file(self, path: Path, *, relative_path: str) -> ParseResult:
        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED:
            return ParseResult(errors=[f"Formato no soportado: {suffix}"])
        supplier = detect_supplier_from_path(relative_path, path.name)
        manufacturer = detect_brand(path.name, supplier)
        if suffix == ".csv":
            result = self._parse_csv(path, relative_path=relative_path, supplier=supplier)
        else:
            result = self._parse_xlsx(path, relative_path=relative_path, supplier=supplier)
        result.supplier = supplier
        result.manufacturer = manufacturer
        return result

    def _parse_csv(self, path: Path, *, relative_path: str, supplier: str | None) -> ParseResult:
        result = ParseResult()
        audit = SheetAudit(sheet_name="CSV", status="empty")
        try:
            text = path.read_text(encoding="utf-8", errors="ignore")
            rows = list(csv.reader(io.StringIO(text)))
            if not rows:
                audit.reason = "archivo_vacio"
                result.sheet_audits.append(audit)
                return result
            header_row_idx, mapping, price_map, stock_map, headers = self._detect_headers(rows)
            audit.header_row = header_row_idx + 1
            audit.columns_mapped = self._columns_mapped_label(mapping, price_map, stock_map, headers)
            if not mapping.get("sku") and not mapping.get("description"):
                audit.status = "ignored"
                audit.reason = "sin_encabezados_reconocibles"
                result.sheet_audits.append(audit)
                result.errors.append("CSV sin encabezados reconocibles")
                return result
            for idx, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
                audit.rows_read += 1
                parsed, reason = self._row_to_product(
                    row, mapping, price_map, stock_map, headers,
                    supplier=supplier, filename=path.name, sheet_name="CSV",
                )
                if parsed:
                    parsed.row_number = idx
                    audit.rows_valid += 1
                    result.rows.append(parsed)
                else:
                    audit.rows_discarded += 1
                    if reason:
                        audit.discard_reasons[reason] = audit.discard_reasons.get(reason, 0) + 1
            audit.status = "indexed" if audit.rows_valid else "empty"
            result.sheet_audits.append(audit)
        except Exception as exc:
            result.errors.append(str(exc))
        return result

    def _parse_xlsx(self, path: Path, *, relative_path: str, supplier: str | None) -> ParseResult:
        result = ParseResult()
        try:
            from openpyxl import load_workbook
        except ImportError:
            result.errors.append("openpyxl no disponible")
            return result
        try:
            wb = load_workbook(path, read_only=False, data_only=True)
        except Exception as exc:
            result.errors.append(f"No se pudo abrir Excel: {exc}")
            return result

        for sheet_name in wb.sheetnames:
            skip = sheet_skip_reason(sheet_name)
            if skip:
                result.sheet_audits.append(SheetAudit(
                    sheet_name=sheet_name, status="ignored", reason=skip,
                ))
                continue
            if not is_commercial_sheet(sheet_name):
                result.sheet_audits.append(SheetAudit(
                    sheet_name=sheet_name, status="ignored", reason="hoja_no_comercial",
                ))
                continue

            ws = wb[sheet_name]
            audit = SheetAudit(sheet_name=sheet_name, status="empty")
            header_row, mapping, price_map, stock_map, headers = self._find_header_row(ws)
            audit.header_row = header_row
            audit.columns_mapped = self._columns_mapped_label(mapping, price_map, stock_map, headers)

            if not mapping.get("sku") and not mapping.get("description"):
                audit.status = "ignored"
                audit.reason = "sin_sku_ni_descripcion"
                result.sheet_audits.append(audit)
                continue

            if not price_map:
                audit.status = "ignored"
                audit.reason = "sin_columna_precio"
                result.sheet_audits.append(audit)
                continue

            sample_valid = 0
            for r in range(header_row + 1, min(header_row + 25, ws.max_row + 1)):
                cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                parsed, _ = self._row_to_product(
                    cells, mapping, price_map, stock_map, headers,
                    supplier=supplier, filename=path.name, sheet_name=sheet_name,
                )
                if parsed:
                    sample_valid += 1
            if sample_valid == 0:
                audit.status = "ignored"
                audit.reason = "sin_filas_validas_muestra"
                result.sheet_audits.append(audit)
                continue

            for r in range(header_row + 1, ws.max_row + 1):
                cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                if not any(c is not None and str(c).strip() for c in cells):
                    continue
                audit.rows_read += 1
                parsed, reason = self._row_to_product(
                    cells, mapping, price_map, stock_map, headers,
                    supplier=supplier, filename=path.name, sheet_name=sheet_name,
                )
                if parsed:
                    parsed.row_number = r
                    audit.rows_valid += 1
                    result.rows.append(parsed)
                else:
                    audit.rows_discarded += 1
                    if reason:
                        audit.discard_reasons[reason] = audit.discard_reasons.get(reason, 0) + 1

            audit.status = "indexed" if audit.rows_valid else "empty"
            result.sheet_audits.append(audit)

        wb.close()
        if not result.rows:
            result.errors.append("No se encontraron filas de productos indexables en hojas comerciales")
        return result

    def _columns_mapped_label(
        self,
        mapping: dict[str, int],
        price_map: dict[str, int],
        stock_map: dict[str, int],
        headers: list[str],
    ) -> dict[str, str]:
        out: dict[str, str] = {}
        for k, idx in mapping.items():
            out[k] = headers[idx] if idx < len(headers) else str(idx)
        for k, idx in price_map.items():
            out[k] = headers[idx] if idx < len(headers) else str(idx)
        for k, idx in stock_map.items():
            out[k] = headers[idx] if idx < len(headers) else str(idx)
        return out

    def _detect_headers(self, rows: list[list[str]]) -> tuple[int, dict, dict, dict, list[str]]:
        best = (0, {}, {}, {}, [])
        best_score = 0
        for idx, row in enumerate(rows[:30]):
            headers = [normalize_header(c) for c in row]
            mapping = self._map_standard(headers)
            price_map = self._map_prices(headers)
            stock_map = self._map_stock(headers)
            score = len(mapping) + len(price_map) * 2
            if "sku" in mapping:
                score += 2
            if score > best_score:
                best_score = score
                best = (idx, mapping, price_map, stock_map, [str(c or "") for c in row])
        return best

    def _find_header_row(self, ws) -> tuple[int, dict, dict, dict, list[str]]:
        best_row = 0
        best_map: dict[str, int] = {}
        best_prices: dict[str, int] = {}
        best_stock: dict[str, int] = {}
        best_headers: list[str] = []
        best_score = 0
        for r in range(1, min(30, ws.max_row + 1)):
            headers_raw = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            headers = [normalize_header(h) for h in headers_raw]
            mapping = self._map_standard(headers)
            price_map = self._map_prices(headers)
            stock_map = self._map_stock(headers)
            score = len(mapping) + len(price_map) * 2
            if "sku" in mapping:
                score += 2
            if "price_final" in price_map:
                score += 3
            if score > best_score:
                best_score = score
                best_row = r
                best_map = mapping
                best_prices = price_map
                best_stock = stock_map
                best_headers = [str(h or "") for h in headers_raw]
        return best_row, best_map, best_prices, best_stock, best_headers

    def _map_standard(self, headers: list[str]) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            for idx, header in enumerate(headers):
                if not header:
                    continue
                for alias in aliases:
                    if header == alias or alias in header:
                        mapping[field] = idx
                        break
                if field in mapping:
                    break
        return mapping

    def _map_prices(self, headers: list[str]) -> dict[str, int]:
        found: dict[str, int] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            for field, aliases in PRICE_COLUMN_RULES:
                if field in found:
                    continue
                for alias in aliases:
                    if header == alias or alias in header:
                        if field == "price_rebate" and "precio" in header:
                            continue
                        found[field] = idx
                        break
        return found

    def _map_stock(self, headers: list[str]) -> dict[str, int]:
        found: dict[str, int] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            for field, aliases in STOCK_COLUMN_RULES:
                if field in found:
                    continue
                for alias in aliases:
                    if header == alias or alias in header:
                        found[field] = idx
                        break
        return found

    def _row_to_product(
        self,
        cells: list[object],
        mapping: dict[str, int],
        price_map: dict[str, int],
        stock_map: dict[str, int],
        headers: list[str],
        *,
        supplier: str | None,
        filename: str,
        sheet_name: str,
    ) -> tuple[ParsedPriceRow | None, str | None]:
        def val(idx: int | None) -> object:
            if idx is None or idx >= len(cells):
                return None
            return cells[idx]

        def field(field_name: str) -> object:
            return val(mapping.get(field_name))

        raw_columns: list[dict[str, Any]] = []
        raw_row: dict[str, Any] = {}
        for i, header in enumerate(headers):
            if not header.strip():
                continue
            cell_val = val(i)
            serial = self._serialize(cell_val)
            raw_row[header] = serial
            raw_columns.append({"column": header, "index": i, "value": serial})

        sku_raw = field("sku")
        sku = str(sku_raw).strip() if sku_raw is not None else None
        if sku and sku.lower() in self.SECTION_HEADERS | {"sku", "product id", "cod. fabricante"}:
            return None, "encabezado_seccion"
        description = field("description")
        desc_text = str(description).strip() if description else ""
        if not sku and not desc_text:
            return None, "fila_vacia"
        if sku and len(sku) < 3 and not desc_text:
            return None, "sku_invalido"

        prices_original: dict[str, Any] = {}
        parsed_prices: dict[str, Decimal | None] = {}
        currency = "USD"
        for price_field, col_idx in price_map.items():
            raw_price = val(col_idx)
            if raw_price is None:
                continue
            p, cur = parse_price(raw_price)
            if p is not None:
                parsed_prices[price_field] = p
                prices_original[price_field] = self._serialize(raw_price)
                currency = cur

        preferred_price: Decimal | None = None
        preferred_field: str | None = None
        for pf in PREFERRED_PRICE_ORDER:
            if pf in parsed_prices and parsed_prices[pf] is not None:
                preferred_price = parsed_prices[pf]
                preferred_field = pf
                break

        if preferred_price is None:
            return None, "sin_precio_valido"

        stock_raw = val(stock_map.get("stock")) if "stock" in stock_map else field("stock")
        transit_raw = val(stock_map.get("transit")) if "transit" in stock_map else None
        stock = parse_int_quantity(stock_raw)
        in_transit = parse_int_quantity(transit_raw)
        stock_parts = []
        if stock_raw is not None:
            stock_parts.append(f"stock={stock_raw}")
        if transit_raw is not None:
            stock_parts.append(f"transito={transit_raw}")
        stock_text = "; ".join(stock_parts) or None

        ram_raw = field("ram")
        storage_raw = field("storage")
        ram_gb = parse_ram_gb(ram_raw)
        storage_gb, storage_type = parse_storage(storage_raw)
        if ram_gb is None and desc_text:
            ram_gb = parse_ram_gb(desc_text)
        if storage_gb is None and desc_text:
            storage_gb, storage_type = parse_storage(desc_text)

        category_raw = field("category")
        category = str(category_raw).strip() if category_raw else None
        classification = classify_product(
            description=desc_text,
            category=category,
            sheet_name=sheet_name,
            preferred_price=preferred_price,
            preferred_price_field=preferred_field,
        )
        product_type = classification.product_type
        brand = detect_brand(filename, desc_text, category, supplier, field("model"))

        stock_source_column = None
        if "stock" in stock_map and stock_map["stock"] < len(headers):
            stock_source_column = headers[stock_map["stock"]].strip()

        search_blob = build_search_blob(
            sku, field("mpn"), field("model"), desc_text, category, brand, supplier,
            json.dumps(raw_row, ensure_ascii=False),
        )

        return ParsedPriceRow(
            sku=sku,
            mpn=str(field("mpn")).strip() if field("mpn") else None,
            model=str(field("model")).strip() if field("model") else None,
            description=desc_text or None,
            category=category,
            product_type=product_type,
            brand=brand,
            processor=str(field("processor")).strip() if field("processor") else None,
            ram_gb=ram_gb,
            storage_gb=storage_gb,
            storage_type=storage_type,
            display=str(field("display")).strip() if field("display") else None,
            operating_system=str(field("os")).strip() if field("os") else None,
            price=preferred_price,
            preferred_price=preferred_price,
            preferred_price_field=preferred_field,
            price_regular=parsed_prices.get("price_regular"),
            price_rebate=parsed_prices.get("price_rebate"),
            price_discount=parsed_prices.get("price_discount"),
            prices_original=prices_original,
            currency=currency,
            stock=stock,
            in_transit=in_transit,
            stock_text_original=stock_text,
            warranty=str(field("warranty")).strip() if field("warranty") else None,
            sheet_name=sheet_name,
            raw_row_json=raw_row,
            raw_columns_json=raw_columns,
            search_blob=search_blob,
            is_commercial=True,
            excluded_from_laptop=classification.excluded_from_laptop,
            is_cotizable=classification.is_cotizable,
            price_review_status=classification.price_review_status,
            classification_label=classification.classification_label,
            stock_source_column=stock_source_column,
        ), None

    @staticmethod
    def _serialize(value: object) -> Any:
        if value is None:
            return None
        if isinstance(value, (str, int, float, bool)):
            return value
        if isinstance(value, Decimal):
            return float(value)
        return str(value)
