"""Investigation Workspace 1.0 MVP — speech acts, ops, Hermes, export flow."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from app.lottery.ai.active_investigation.hermes_decision_engine import HermesDecisionEngine
from app.lottery.ai.active_investigation.session import ActiveInvestigationSession
from app.lottery.ai.conversation_state import ConversationState
from app.lottery.ai.investigation_workspace.export_xlsx import export_asset_xlsx
from app.lottery.ai.investigation_workspace.handler import (
    execute_workspace_action,
    resolve_export_file,
)
from app.lottery.ai.investigation_workspace.materialize import materialize_same_day_table
from app.lottery.ai.investigation_workspace.operations import set_filters, set_sort
from app.lottery.ai.investigation_workspace.schemas import AssetFilters, AssetSort
from app.lottery.ai.investigation_workspace.speech_acts import WorkspaceSpeechActDetector
from app.lottery.ai.investigation_workspace.store import save_asset


def _sample_items() -> list[dict]:
    return [
        {
            "date": "2024-01-15",
            "appearances": [
                {"number": "35", "lottery": "Quiniela Loteka", "position": 1},
                {"number": "14", "lottery": "Quiniela Loteka", "position": 2},
            ],
        },
        {
            "date": "2023-06-01",
            "appearances": [
                {"number": "35", "lottery": "Quiniela Leidsa", "position": 1},
                {"number": "14", "lottery": "Quiniela Real", "position": 3},
            ],
        },
        {
            "date": "2025-03-10",
            "appearances": [
                {"number": "35", "lottery": "Lotería Nacional", "position": 2},
                {"number": "14", "lottery": "Quiniela Loteka", "position": 1},
            ],
        },
    ]


def test_speech_acts_mvp_phrases():
    phrases = [
        ("Muéstrame esos resultados", "show_results"),
        ("Filtra solo Loteka", "filter_results"),
        ("Ordénalos por fecha", "sort_results"),
        ("Exporta a Excel", "export_results"),
    ]
    for text, action in phrases:
        d = WorkspaceSpeechActDetector.detect(
            text, has_active_asset=True, has_active_investigation=True
        )
        assert d is not None, text
        assert d.action == action, text
    filt = WorkspaceSpeechActDetector.detect(
        "Filtra solo Loteka", has_active_asset=True, has_active_investigation=True
    )
    assert filt is not None
    assert filt.filters.lottery == "Loteka"


def test_routing_residuals_never_workspace_without_unequivocal_asset_act():
    """Cert200 residuals: Analyst 2.1 factual turns must not become asset_action."""
    residual = [
        "¿Solo en 2026?",
        "Últimas 3 del 35 en Nacional.",
        "Ahora en todas las posiciones.",
        "Ahora en todas las loterías.",
        "En Nacional.",
        "¿Y el 44?",
    ]
    for text in residual:
        # Even with a stale asset / active investigation — factual phrasing wins.
        assert (
            WorkspaceSpeechActDetector.detect(
                text, has_active_asset=True, has_active_investigation=True
            )
            is None
        ), text
        assert (
            WorkspaceSpeechActDetector.detect(
                text, has_active_asset=False, has_active_investigation=True
            )
            is None
        ), text


def test_routing_abbreviated_filter_requires_active_asset():
    assert (
        WorkspaceSpeechActDetector.detect(
            "Solo Loteka", has_active_asset=False, has_active_investigation=True
        )
        is None
    )
    assert (
        WorkspaceSpeechActDetector.detect(
            "Todas las loterías", has_active_asset=False, has_active_investigation=True
        )
        is None
    )
    d = WorkspaceSpeechActDetector.detect(
        "Solo Loteka", has_active_asset=True, has_active_investigation=True
    )
    assert d is not None and d.action == "filter_results"
    assert d.filters.lottery == "Loteka"
    clear = WorkspaceSpeechActDetector.detect(
        "Todas las loterías", has_active_asset=True, has_active_investigation=True
    )
    assert clear is not None and clear.reason_code == "clear_filters"


def test_hermes_residuals_not_asset_action():
    state = ConversationState(
        active_numbers=["54", "94"],
        active_pair=["54", "94"],
        active_relation="same_day",
    )
    inv = ActiveInvestigationSession(
        subjects=["54", "94"], relation="same_day", metric="same_day"
    )
    asset = materialize_same_day_table(
        items=_sample_items(), subjects=["54", "94"], investigation_id=inv.investigation_id
    )
    save_asset(state, asset)
    for text in (
        "¿Solo en 2026?",
        "Últimas 3 del 35 en Nacional.",
        "Ahora en todas las posiciones.",
        "Ahora en todas las loterías.",
        "En Nacional.",
    ):
        d = HermesDecisionEngine.decide(text, state=state, investigation=inv)
        assert d.turn_type != "asset_action", text
        assert d.workspace_action is None, text


def test_hermes_positive_filter_sort_export_with_asset():
    state = ConversationState(
        active_numbers=["35", "14"],
        active_pair=["35", "14"],
        active_relation="same_day",
    )
    inv = ActiveInvestigationSession(
        subjects=["35", "14"], relation="same_day", metric="same_day"
    )
    asset = materialize_same_day_table(
        items=_sample_items(), subjects=["35", "14"], investigation_id=inv.investigation_id
    )
    save_asset(state, asset)
    for text, action in (
        ("Filtra solo Loteka", "filter_results"),
        ("Ordénalos por fecha", "sort_results"),
        ("Exporta a Excel", "export_results"),
        ("En esa tabla, solo Nacional", "filter_results"),
    ):
        d = HermesDecisionEngine.decide(text, state=state, investigation=inv)
        assert d.turn_type == "asset_action", text
        assert d.workspace_action is not None
        assert d.workspace_action["action"] == action, text


def test_hermes_prefers_asset_action_over_default_research():
    state = ConversationState(
        active_numbers=["35", "14"],
        active_pair=["35", "14"],
        active_relation="same_day",
    )
    inv = ActiveInvestigationSession(
        subjects=["35", "14"],
        relation="same_day",
        metric="same_day",
    )
    asset = materialize_same_day_table(
        items=_sample_items(), subjects=["35", "14"], investigation_id=inv.investigation_id
    )
    save_asset(state, asset)
    d = HermesDecisionEngine.decide(
        "Muéstrame esos resultados", state=state, investigation=inv
    )
    assert d.turn_type == "asset_action"
    assert d.requires_research is False
    assert d.reasoning_mode == "skip"
    assert d.workspace_action is not None
    assert d.workspace_action["action"] == "show_results"


def test_filter_sort_recompute():
    asset = materialize_same_day_table(items=_sample_items(), subjects=["35", "14"])
    assert asset.row_count == 3
    asset = set_filters(asset, AssetFilters(lottery="Loteka"))
    assert asset.row_count == 2
    asset = set_sort(asset, AssetSort(field="fecha", direction="asc"))
    dates = [r["fecha"] for r in asset.rows]
    assert dates == sorted(dates)


@pytest.mark.asyncio
async def test_execute_full_flow_mock_query(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "app.config.settings.lottery_exports_path", str(tmp_path)
    )
    state = ConversationState(
        active_numbers=["35", "14"],
        active_pair=["35", "14"],
        active_relation="same_day",
    )
    inv = ActiveInvestigationSession(subjects=["35", "14"], relation="same_day")
    query = SimpleNamespace(
        same_day_number_coincidences=AsyncMock(
            return_value={"items": _sample_items(), "total": 3}
        )
    )

    show = WorkspaceSpeechActDetector.detect(
        "Muéstrame esos resultados",
        has_active_asset=False,
        has_active_investigation=True,
    )
    assert show is not None
    r1 = await execute_workspace_action(
        state, show, query_service=query, investigation=inv, conversation_id="c1"
    )
    assert r1["tool_payload"]["ok"] is True
    assert "fecha" in (r1["content"] or "")
    assert (r1["structured_content"] or {}).get("type") == "investigation_workspace_table"
    assert r1["structured_content"]["asset"]["row_count"] == 3

    filt = WorkspaceSpeechActDetector.detect(
        "Filtra solo Loteka", has_active_asset=True, has_active_investigation=True
    )
    r2 = await execute_workspace_action(
        state, filt, query_service=query, investigation=inv
    )
    assert r2["structured_content"]["asset"]["row_count"] == 2

    sort = WorkspaceSpeechActDetector.detect(
        "Ordénalos por fecha", has_active_asset=True, has_active_investigation=True
    )
    r3 = await execute_workspace_action(
        state, sort, query_service=query, investigation=inv
    )
    rows = r3["structured_content"]["asset"]["rows"]
    # page may truncate; check asset order via store
    from app.lottery.ai.investigation_workspace.store import get_active_asset

    asset = get_active_asset(state)
    assert asset is not None
    dates = [r["fecha"] for r in asset.rows]
    assert dates == sorted(dates, reverse=True)

    exp = WorkspaceSpeechActDetector.detect(
        "Exporta a Excel", has_active_asset=True, has_active_investigation=True
    )
    r4 = await execute_workspace_action(
        state, exp, query_service=query, investigation=inv
    )
    assert r4["structured_content"]["type"] == "investigation_workspace_export"
    storage = r4["structured_content"]["storage_name"]
    path = resolve_export_file(storage)
    assert path is not None
    assert path.is_file()
    assert path.stat().st_size > 100
    # openpyxl sheets
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert "Resumen" in wb.sheetnames
    assert "Resultados" in wb.sheetnames
    assert wb["Resultados"].max_row == 1 + 2  # header + 2 Loteka rows


def test_export_xlsx_writes_file(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "app.config.settings.lottery_exports_path", str(tmp_path)
    )
    asset = materialize_same_day_table(items=_sample_items(), subjects=["35", "14"])
    path, filename, content = export_asset_xlsx(asset)
    assert path.exists()
    assert filename.endswith(".xlsx")
    assert content[:2] == b"PK"
