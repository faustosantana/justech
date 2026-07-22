"""Relleno de plantillas SNCC oficiales — corchetes y etiquetas de tabla."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.services.document_autofill.docx_template_writer import fill_docx_bytes

OFFICIAL_F042 = Path("/var/jaios/documents/m365_template_cache/709c0aa6-1281-4e1a-b959-2a4c74de106a.docx")


@pytest.mark.skipif(not OFFICIAL_F042.is_file(), reason="plantilla oficial no en caché")
def test_sncc_f042_fills_corporate_body_fields():
    template_bytes = OFFICIAL_F042.read_bytes()
    values = {
        "razon_social": "Justech SRL",
        "rnc": "131828061",
        "direccion": "Calle Francisco Segura Sandoval #157, Reparto Alma Rosa, SDE",
        "representante_autorizado": "Fausto Ramón Santana Jiménez",
        "telefono": "+1-809-555-0100",
        "correo": "fausto@justech.do",
        "proceso_dgcp": "DGII-CCC-PEEX-2026-0005",
        "fecha": "24/06/2026",
        "entidad_contratante": "Dirección General Impuestos Internos",
    }
    alias_values = {
        "Fecha de emisión del documento": "24/06/2026",
        "No. del Expediente de Compras ": "DGII-CCC-PEEX-2026-0005",
        "Nombre de la Institución": "Dirección General Impuestos Internos",
    }
    filled = fill_docx_bytes(template_bytes, values, alias_values=alias_values)
    doc = Document(BytesIO(filled))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    assert "Justech SRL" in text
    assert "131828061" in text
    assert "Fausto Ramón Santana Jiménez" in text
    assert "Alma Rosa" in text
    assert "[indicar el nombre jurídico del Oferente]" not in text
